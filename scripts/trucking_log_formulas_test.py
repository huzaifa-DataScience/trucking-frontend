#!/usr/bin/env python3
"""
Executable test cases for 2025-2026 Trucking Log.xlsx formulas.

Run from repo root:
  python3 scripts/trucking_log_formulas_test.py

Implements the same costing math as Trucking Record columns AM / AN / AO
and asserts against synthetic cases + cached workbook values.
"""

from __future__ import annotations

import sys
import unittest
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
XLSX = ROOT / "2025-2026 Trucking Log.xlsx"


# ── Pure formula ports (no Excel needed for unit cases) ─────────────────────


def hauler_key(company: str, category: str) -> str:
    """Haulers!A = IF(B="","Z", CONCATENATE(B," - ",C))"""
    if company == "":
        return "Z"
    return f"{company} - {category}"


def disposal_key(company: str, column1: str, category: str) -> str:
    """Disposal!A = IF(B="","X", CONCATENATE(B," - ",I," - ",J))"""
    if company == "":
        return "X"
    return f"{company} - {column1} - {category}"


def trucking_cost(
    hauler_selected: str,
    rate_per_load: float | None,
    rate_per_hr: float | None,
    hours: float | None,
) -> float:
    """
    AM = IF(D="-----", 0, VLOOKUP(...,4) + VLOOKUP(...,5)*E)
    Blank VLOOKUP numerics → 0; blank hours → 0.
    """
    if hauler_selected == "-----":
        return 0.0
    load = float(rate_per_load or 0)
    hr = float(rate_per_hr or 0)
    hrs = float(hours or 0)
    return load + hr * hrs


def billable_tons(weight: float | None, included_tons: float | None) -> float:
    """
    IF(AND(G>0, (G-M)>0), 0, M-G)
    Blank G treated as 0.
    """
    m = float(weight or 0)
    g = float(included_tons or 0)
    if g > 0 and (g - m) > 0:
        return 0.0
    return m - g


def disposal_cost(
    destination_selected: str,
    cost_per_load: float | None,
    cost_per_ton: float | None,
    cost_per_cy: float | None,
    cy_truck_size: float | None,
    included_tons: float | None,
    weight_tons: float | None,
) -> float:
    """AN disposal formula (load + ton*billable + cy*size)."""
    if destination_selected == "-----":
        return 0.0
    load = float(cost_per_load or 0)
    ton = float(cost_per_ton or 0)
    cy = float(cost_per_cy or 0)
    size = float(cy_truck_size or 0)
    billable = billable_tons(weight_tons, included_tons)
    return load + ton * billable + cy * size


def total_cost(trucking: float, disposal: float) -> float:
    """AO = SUM(AM:AN)"""
    return trucking + disposal


def lbs_to_short_tons(lbs: float) -> float:
    return lbs / 2000.0


def lbs_to_long_tons(lbs: float) -> float:
    return lbs / 2240.0


# ── Unit tests ───────────────────────────────────────────────────────────────


class TestLookupKeys(unittest.TestCase):
    def test_key_hauler_a_and_s_30yd(self):
        self.assertEqual(
            hauler_key("A&S Refuse", "Hauler  - 30 YD can C&D"),
            "A&S Refuse - Hauler  - 30 YD can C&D",
        )

    def test_key_hauler_blank_company(self):
        self.assertEqual(hauler_key("", "Hauler"), "Z")

    def test_key_disposal_a_and_s_cd(self):
        self.assertEqual(
            disposal_key("A&S Refuse", "Disposal", "C&D "),
            "A&S Refuse - Disposal - C&D ",
        )

    def test_key_disposal_blank_company(self):
        self.assertEqual(disposal_key("", "Export", "Dirt"), "X")


class TestTruckingCost(unittest.TestCase):
    def test_dash_hauler_is_zero(self):
        self.assertEqual(trucking_cost("-----", 550, 100, 3), 0.0)

    def test_per_load_only_ignores_hours(self):
        self.assertEqual(trucking_cost("A&S…", 550, None, None), 550.0)
        self.assertEqual(trucking_cost("A&S…", 550, None, 3), 550.0)

    def test_hourly_one_hour(self):
        self.assertEqual(trucking_cost("Hammer…", None, 85, 1), 85.0)

    def test_hourly_fraction(self):
        self.assertAlmostEqual(trucking_cost("Hammer…", None, 85, 0.95), 80.75)

    def test_hourly_blank_hours_is_zero(self):
        self.assertEqual(trucking_cost("Hammer…", None, 85, None), 0.0)
        self.assertEqual(trucking_cost("Hammer…", None, 85, 0), 0.0)

    def test_synthetic_load_plus_hourly(self):
        self.assertEqual(trucking_cost("Both", 200, 100, 2), 400.0)


class TestBillableTons(unittest.TestCase):
    def test_no_included_charges_full_weight(self):
        self.assertEqual(billable_tons(9.84, 0), 9.84)
        self.assertEqual(billable_tons(9.84, None), 9.84)

    def test_under_included_is_zero(self):
        self.assertEqual(billable_tons(3, 5), 0.0)

    def test_exact_included_is_zero(self):
        # (G-M)>0 is false when equal → falls through to M-G = 0
        self.assertEqual(billable_tons(5, 5), 0.0)

    def test_over_included_charges_delta(self):
        self.assertEqual(billable_tons(8, 5), 3.0)
        self.assertEqual(billable_tons(10, 5), 5.0)


class TestDisposalCost(unittest.TestCase):
    def test_dash_destination_is_zero(self):
        self.assertEqual(disposal_cost("-----", 500, 67, None, None, 0, 10), 0.0)

    def test_load_only(self):
        self.assertEqual(
            disposal_cost("A&S…", 500, None, None, None, 0, 10.59), 500.0
        )

    def test_ton_only_row5(self):
        self.assertAlmostEqual(
            disposal_cost("Cumberland…", None, 67, None, None, 0, 9.84), 659.28
        )

    def test_included_tons_under(self):
        self.assertEqual(disposal_cost("Recycle…", None, 65, None, None, 5, 3), 0.0)

    def test_included_tons_over(self):
        self.assertEqual(disposal_cost("Recycle…", None, 65, None, None, 5, 8), 195.0)

    def test_load_plus_ton(self):
        self.assertEqual(disposal_cost("Mix…", 100, 50, None, None, 0, 2), 200.0)

    def test_cy_path(self):
        self.assertEqual(
            disposal_cost("CY…", None, None, 10, 20, 0, 99), 200.0
        )


class TestWeightConversion(unittest.TestCase):
    def test_short_tons(self):
        self.assertAlmostEqual(lbs_to_short_tons(3020), 1.51)

    def test_gross_minus_tare(self):
        self.assertAlmostEqual(lbs_to_short_tons(13360 - 11100), 1.13)

    def test_long_tons(self):
        self.assertAlmostEqual(lbs_to_long_tons(1360), 1360 / 2240)

    def test_multi_ticket_sum(self):
        self.assertAlmostEqual(lbs_to_short_tons(746 + 1580 + 953), 1.6395)


class TestTotalCost(unittest.TestCase):
    def test_row6_style(self):
        self.assertEqual(total_cost(550, 500), 1050)

    def test_row5_style(self):
        self.assertAlmostEqual(total_cost(630, 659.28), 1289.28)

    def test_zero(self):
        self.assertEqual(total_cost(0, 0), 0)


class TestTicketBookRange(unittest.TestCase):
    def test_fifty_ticket_book(self):
        start = 162001
        end = start + 49
        self.assertEqual(end, 162050)
        self.assertEqual(end + 1, 162051)


# ── Integration vs Excel cached values ───────────────────────────────────────


@unittest.skipUnless(XLSX.exists(), "workbook not present")
class TestAgainstWorkbookCache(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        import openpyxl

        cls.wb = openpyxl.load_workbook(XLSX, data_only=False)
        cls.wbv = openpyxl.load_workbook(XLSX, data_only=True)
        cls.rec = cls.wb["Trucking Record"]
        cls.recv = cls.wbv["Trucking Record"]

    def _cached(self, row: int, col: int):
        return self.recv.cell(row, col).value

    def test_int_row5_cumberland_asbestos(self):
        # AM=39, AN=40, AO=41 (1-based)
        self.assertEqual(self._cached(5, 39), 630)
        self.assertAlmostEqual(float(self._cached(5, 40)), 659.28)
        self.assertAlmostEqual(float(self._cached(5, 41)), 1289.28)
        # Recompute from visible rate card
        am = trucking_cost("R L Mc Nair…", 630, None, None)
        an = disposal_cost("Cumberland…", None, 67, None, None, 0, 9.84)
        self.assertEqual(am, 630)
        self.assertAlmostEqual(an, 659.28)
        self.assertAlmostEqual(total_cost(am, an), 1289.28)

    def test_int_row6_as_cd_load(self):
        self.assertEqual(self._cached(6, 39), 550)
        self.assertEqual(self._cached(6, 40), 500)
        self.assertEqual(self._cached(6, 41), 1050)

    def test_int_row23_hourly(self):
        self.assertEqual(self.rec.cell(23, 5).value, 1)  # hours
        self.assertEqual(self._cached(23, 39), 85)

    def test_int_row30_hourly_fraction(self):
        self.assertEqual(self.rec.cell(30, 5).value, 0.95)
        self.assertAlmostEqual(float(self._cached(30, 39)), 80.75)

    def test_int_row4_documents_an_override(self):
        """Row 4 AN is hardcoded 53.6; Disposal rate card for that key is $500/load."""
        self.assertEqual(self._cached(4, 39), 550)
        self.assertAlmostEqual(float(self._cached(4, 40)), 53.6)
        # Formula path would be 500 — override is intentional data quirk
        formula_an = disposal_cost("A&S Refuse - Disposal - C&D ", 500, None, None, None, 0, 10.59)
        self.assertEqual(formula_an, 500)
        self.assertNotAlmostEqual(float(self._cached(4, 40)), formula_an)

    def test_weight_formulas_cached(self):
        self.assertAlmostEqual(float(self._cached(53, 13)), 1.51)  # 3020/2000
        self.assertAlmostEqual(float(self._cached(41, 13)), 1.13)  # (13360-11100)/2000


def main() -> int:
    loader = unittest.defaultTestLoader
    suite = unittest.TestSuite()
    suite.addTests(loader.loadTestsFromModule(sys.modules[__name__]))
    result = unittest.TextTestRunner(verbosity=2).run(suite)
    return 0 if result.wasSuccessful() else 1


if __name__ == "__main__":
    raise SystemExit(main())
